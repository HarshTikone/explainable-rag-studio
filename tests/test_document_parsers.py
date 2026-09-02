from pathlib import Path

import fitz
import pytest
from docx import Document as WordDocument

from backend.document_parsers import DocumentParseError, _ocr_page, markdown_table, parse_document, stable_document_id


def test_markdown_parser_preserves_frontmatter_headings_tables_and_unicode(tmp_path):
    path = tmp_path / "Guide.md"
    path.write_text(
        "---\ntitle: Café Guide\nversion: 2.1\n---\n# Install\nUse TS-999 safely.\n\n## Limits\n| Key | Value |\n| --- | --- |\n| max | 10 |\n",
        encoding="utf-8",
    )
    parsed = parse_document(str(path))
    assert parsed.document.title == "Café Guide"
    assert parsed.version.source_version == "2.1"
    assert parsed.blocks[-1].block_type == "table"
    assert parsed.blocks[-1].heading_path == ("Install", "Limits")
    assert parsed.document.document_id == stable_document_id("Guide.md")


def test_html_and_docx_parsers_keep_structure_and_tables(tmp_path):
    html = tmp_path / "manual.html"
    html.write_text("<html><head><title>Manual</title></head><body><h1>Setup</h1><p>Install it.</p><table><tr><th>ID</th><th>State</th></tr><tr><td>A-1</td><td>ready</td></tr></table></body></html>", encoding="utf-8")
    parsed_html = parse_document(str(html))
    assert parsed_html.document.title == "Manual"
    assert {block.block_type for block in parsed_html.blocks} >= {"heading", "paragraph", "table"}

    docx_path = tmp_path / "manual.docx"
    document = WordDocument()
    document.add_heading("Operations", level=1)
    document.add_paragraph("Rotate credentials.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "ID", "State"
    table.cell(1, 0).text, table.cell(1, 1).text = "AU-1", "active"
    document.save(docx_path)
    parsed_docx = parse_document(str(docx_path))
    assert parsed_docx.blocks[0].heading_path == ("Operations",)
    assert parsed_docx.blocks[-1].block_type == "table"


def test_pdf_parser_extracts_text_and_reports_capabilities(tmp_path):
    path = tmp_path / "sample.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Current procedure AU-4012 uses a 45 minute token lifetime.")
    document.save(path)
    document.close()
    parsed = parse_document(str(path))
    assert "AU-4012" in parsed.blocks[0].text
    assert parsed.blocks[0].page == 1
    assert parsed.capabilities["tables"]


def test_empty_unsupported_and_ocr_required_errors(tmp_path, monkeypatch):
    empty = tmp_path / "empty.txt"
    empty.write_text("  ", encoding="utf-8")
    with pytest.raises(DocumentParseError, match="No searchable content") as error:
        parse_document(str(empty))
    assert error.value.code == "EMPTY_DOCUMENT"
    unsupported = tmp_path / "data.csv"
    unsupported.write_text("a,b", encoding="utf-8")
    with pytest.raises(DocumentParseError) as error:
        parse_document(str(unsupported))
    assert error.value.code == "UNSUPPORTED_FORMAT"

    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"fake")
    monkeypatch.setattr(
        "backend.document_parsers._parse_pdf",
        lambda *_args: ([], [{"code": "OCR_UNAVAILABLE", "message": "missing"}], {"ocr": False}),
    )
    with pytest.raises(DocumentParseError) as error:
        parse_document(str(pdf))
    assert error.value.code == "OCR_REQUIRED"


def test_markdown_table_normalizes_ragged_rows():
    value = markdown_table([["Name", "Value"], ["α", 1], ["only-one"]])
    assert "| Name | Value |" in value
    assert "| only-one |  |" in value
    assert markdown_table([[]]) == ""


def test_ocr_adapter_supports_executable_and_wraps_failures(monkeypatch):
    class Pixmap:
        width, height, samples = 1, 1, b"\xff\xff\xff"

    class Page:
        def get_pixmap(self, **options):
            assert options == {"dpi": 200, "alpha": False}
            return Pixmap()

    monkeypatch.setattr("pytesseract.image_to_string", lambda _image: "recognized")
    assert _ocr_page(Page(), "custom-tesseract") == "recognized"
    monkeypatch.setattr("pytesseract.image_to_string", lambda _image: (_ for _ in ()).throw(RuntimeError("bad OCR")))
    with pytest.raises(DocumentParseError) as error:
        _ocr_page(Page(), None)
    assert error.value.code == "OCR_UNAVAILABLE"


def test_scanned_pdf_uses_ocr_result(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()
    monkeypatch.setattr("backend.document_parsers._ocr_page", lambda _page, _executable: "OCR text with enough searchable characters AU-7")
    parsed = parse_document(str(path))
    assert parsed.blocks[0].block_type == "ocr"
    assert "AU-7" in parsed.blocks[0].text
