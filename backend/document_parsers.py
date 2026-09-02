"""Structure-aware parsers with optional OCR and graceful capability checks."""
from __future__ import annotations

import hashlib
import mimetypes
import os
import re
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from .ingestion_models import Document, DocumentVersion, ParseResult, ParsedBlock

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".html", ".htm", ".txt"}
PARSER_VERSION = "1.0"


class DocumentParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def normalize_text(text: str) -> str:
    return " ".join((text or "").replace("\x00", " ").split())


def source_sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def logical_source_identity(source_name: str) -> str:
    return re.sub(r"[^a-z0-9._-]+", "-", Path(source_name).name.casefold()).strip("-")


def stable_document_id(source_name: str, organization_id: str = "org_public") -> str:
    logical = logical_source_identity(source_name)
    identity = logical if organization_id == "org_public" else f"{organization_id}\x1f{logical}"
    return "doc_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:16]


def _block(block_type: str, text: str, page: int | None, headings: Sequence[str], ordinal: int) -> ParsedBlock | None:
    cleaned = normalize_text(text)
    if not cleaned:
        return None
    return ParsedBlock(f"b{ordinal:06d}", block_type, cleaned, page, tuple(headings), ordinal)


def markdown_table(rows: Iterable[Iterable[object]]) -> str:
    normalized = [[normalize_text(str(cell or "")) for cell in row] for row in rows]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return ""
    width = max(len(row) for row in normalized)
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    header = normalized[0]
    return "\n".join([
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
        *("| " + " | ".join(row) + " |" for row in normalized[1:]),
    ])


def _parse_markdown(text: str) -> Tuple[List[ParsedBlock], Dict[str, str]]:
    lines = text.replace("\r\n", "\n").split("\n")
    metadata: Dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        try:
            end = lines.index("---", 1)
            for line in lines[1:end]:
                if ":" in line:
                    key, value = line.split(":", 1)
                    metadata[key.strip().casefold()] = value.strip().strip("\"'")
            lines = lines[end + 1:]
        except ValueError:
            pass
    headings: List[str] = []
    blocks: List[ParsedBlock] = []
    buffer: List[str] = []

    def flush(block_type: str = "paragraph"):
        if buffer:
            candidate = _block(block_type, "\n".join(buffer), 1, headings, len(blocks) + 1)
            if candidate:
                blocks.append(candidate)
            buffer.clear()

    for line in lines:
        match = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if match:
            flush()
            level, title = len(match.group(1)), normalize_text(match.group(2))
            headings[:] = headings[:level - 1]
            headings.append(title)
            candidate = _block("heading", title, 1, headings, len(blocks) + 1)
            if candidate:
                blocks.append(candidate)
        elif not line.strip():
            flush("table" if buffer and all("|" in value for value in buffer) else "paragraph")
        else:
            buffer.append(line)
    flush("table" if buffer and all("|" in value for value in buffer) else "paragraph")
    return blocks, metadata


def _parse_html(path: str) -> Tuple[List[ParsedBlock], Dict[str, str]]:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise DocumentParseError("PARSER_UNAVAILABLE", "BeautifulSoup is required for HTML ingestion.") from exc
    soup = BeautifulSoup(Path(path).read_text(encoding="utf-8"), "html.parser")
    headings: List[str] = []
    blocks: List[ParsedBlock] = []
    for node in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"]):
        if node.name.startswith("h"):
            level = int(node.name[1])
            title = normalize_text(node.get_text(" "))
            headings[:] = headings[:level - 1]
            headings.append(title)
            kind, value = "heading", title
        elif node.name == "table":
            kind = "table"
            value = markdown_table([[cell.get_text(" ") for cell in row.find_all(["th", "td"])] for row in node.find_all("tr")])
        else:
            kind, value = ("list" if node.name == "li" else "paragraph"), node.get_text(" ")
        candidate = _block(kind, value, 1, headings, len(blocks) + 1)
        if candidate:
            blocks.append(candidate)
    title = normalize_text(soup.title.get_text(" ")) if soup.title else ""
    return blocks, {"title": title} if title else {}


def _parse_docx(path: str) -> Tuple[List[ParsedBlock], Dict[str, str]]:
    try:
        from docx import Document as WordDocument
    except ImportError as exc:
        raise DocumentParseError("PARSER_UNAVAILABLE", "python-docx is required for DOCX ingestion.") from exc
    document = WordDocument(path)
    headings: List[str] = []
    blocks: List[ParsedBlock] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").casefold()
        match = re.match(r"heading\s+(\d+)", style)
        if match:
            level = max(1, min(6, int(match.group(1))))
            headings[:] = headings[:level - 1]
            headings.append(text)
            kind = "heading"
        else:
            kind = "paragraph"
        blocks.append(_block(kind, text, 1, headings, len(blocks) + 1))
    for table in document.tables:
        value = markdown_table([[cell.text for cell in row.cells] for row in table.rows])
        candidate = _block("table", value, 1, headings, len(blocks) + 1)
        if candidate:
            blocks.append(candidate)
    return blocks, {}


def _ocr_page(page, executable: str | None) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise DocumentParseError("OCR_UNAVAILABLE", "pytesseract and Pillow are required for OCR.") from exc
    if executable:
        pytesseract.pytesseract.tesseract_cmd = executable
    pixmap = page.get_pixmap(dpi=200, alpha=False)
    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
    try:
        return pytesseract.image_to_string(image)
    except Exception as exc:
        raise DocumentParseError("OCR_UNAVAILABLE", f"Tesseract OCR failed: {exc}") from exc


def _parse_pdf(path: str, ocr_executable: str | None) -> Tuple[List[ParsedBlock], List[Dict[str, str]], Dict[str, bool]]:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentParseError("PARSER_UNAVAILABLE", "PyMuPDF is required for PDF ingestion.") from exc
    warnings: List[Dict[str, str]] = []
    blocks: List[ParsedBlock] = []
    ocr_available = True
    with fitz.open(path) as pdf:
        for page_number, page in enumerate(pdf, 1):
            page_text = normalize_text(page.get_text("text"))
            if len(re.sub(r"\W", "", page_text)) < 20:
                try:
                    page_text = normalize_text(_ocr_page(page, ocr_executable))
                except DocumentParseError as exc:
                    ocr_available = False
                    warnings.append({"code": exc.code, "message": f"Page {page_number}: {exc}"})
            candidate = _block("ocr" if not normalize_text(page.get_text("text")) and page_text else "paragraph", page_text, page_number, (), len(blocks) + 1)
            if candidate:
                blocks.append(candidate)
    table_available = True
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                for table in page.extract_tables() or []:
                    candidate = _block("table", markdown_table(table), page_number, (), len(blocks) + 1)
                    if candidate:
                        blocks.append(candidate)
    except (ImportError, Exception) as exc:
        table_available = False
        warnings.append({"code": "TABLE_EXTRACTION_UNAVAILABLE", "message": str(exc)})
    return blocks, warnings, {"ocr": ocr_available, "tables": table_available}


def parse_document(path: str, *, source_name: str | None = None, source_version: str | None = None,
                   ocr_executable: str | None = None, organization_id: str = "org_public",
                   created_by: str = "", trust_state: str = "public") -> ParseResult:
    file_path = Path(path)
    suffix = file_path.suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentParseError("UNSUPPORTED_FORMAT", f"Unsupported document type: {suffix or 'unknown'}")
    name = Path(source_name or file_path.name).name
    checksum = source_sha256(path)
    metadata: Dict[str, str] = {}
    warnings: List[Dict[str, str]] = []
    capabilities = {"ocr": False, "tables": False}
    if suffix == ".pdf":
        blocks, warnings, capabilities = _parse_pdf(path, ocr_executable)
    elif suffix == ".docx":
        blocks, metadata = _parse_docx(path)
        capabilities["tables"] = True
    elif suffix in {".html", ".htm"}:
        blocks, metadata = _parse_html(path)
        capabilities["tables"] = True
    elif suffix in {".md", ".markdown"}:
        blocks, metadata = _parse_markdown(file_path.read_text(encoding="utf-8"))
    else:
        value = file_path.read_text(encoding="utf-8").replace("\x00", " ")
        candidate = _block("paragraph", value, 1, (), 1)
        blocks = [candidate] if candidate else []
    if not blocks:
        code = "OCR_REQUIRED" if suffix == ".pdf" and any(warning["code"] == "OCR_UNAVAILABLE" for warning in warnings) else "EMPTY_DOCUMENT"
        raise DocumentParseError(code, "No searchable content could be extracted.")
    title = metadata.get("title") or next((block.text for block in blocks if block.block_type == "heading"), file_path.stem)
    document_id = stable_document_id(name, organization_id)
    document = Document(document_id, logical_source_identity(name), title, mimetypes.guess_type(name)[0] or "application/octet-stream",
                        organization_id, created_by, created_by, trust_state)
    inferred_version = metadata.get("version")
    if not inferred_version:
        preview = " ".join(block.text for block in blocks[:3])
        status_match = re.search(r"\bstatus\s*:\s*([a-z0-9._-]+)", preview, flags=re.IGNORECASE)
        inferred_version = status_match.group(1).casefold().rstrip(".,;:") if status_match else None
    if not inferred_version:
        name_match = re.search(r"_(current|legacy|incident)(?:\.[^.]+)?$", name, flags=re.IGNORECASE)
        inferred_version = name_match.group(1).casefold() if name_match else "general"
    version_identity = checksum if organization_id == "org_public" else hashlib.sha256(f"{organization_id}\x1f{checksum}".encode()).hexdigest()
    version = DocumentVersion("ver_" + version_identity[:20], document_id, checksum, source_version or inferred_version, os.path.getsize(path), organization_id)
    return ParseResult(document, version, blocks, warnings, capabilities)
